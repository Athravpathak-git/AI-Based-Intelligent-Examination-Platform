import io
import json
import pytest
import openpyxl
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

def test_csv_preview_and_confirm(client, examiner_auth_headers):
    # CSV with 1 valid MCQ, 1 valid Multi-Select, 1 Invalid, and 1 Duplicate
    # First seed a duplicate into database
    client.post(
        "/api/questions",
        json={
            "subject": "Chemistry",
            "question_text": "What is the chemical symbol for Gold?",
            "question_type": "MCQ",
            "difficulty": "EASY",
            "marks": 2.0,
            "negative_marks": 0.5,
            "options": [
                {"option_text": "Au", "is_correct": True, "option_order": 0},
                {"option_text": "Ag", "is_correct": False, "option_order": 1}
            ]
        },
        headers=examiner_auth_headers
    )

    csv_content = """subject,topic,question_text,question_type,difficulty,marks,negative_marks,option_1,option_1_correct,option_2,option_2_correct,option_3,option_3_correct,expected_answer,explanation
Chemistry,Elements,What is the chemical symbol for Water?,MCQ,EASY,1.0,0.25,H2O,TRUE,CO2,FALSE,O2,FALSE,,Water is H2O
Chemistry,Elements,Which of the following are noble gases?,MULTI_SELECT,MEDIUM,3.0,1.0,Helium,TRUE,Neon,TRUE,Oxygen,FALSE,,Noble gases have full valence shells
,NoSubject,Invalid question without subject,MCQ,EASY,1.0,0.0,A,TRUE,B,FALSE,,,
Chemistry,Elements,What is the chemical symbol for Gold?,MCQ,EASY,2.0,0.5,Au,TRUE,Ag,FALSE,,,Duplicate in DB
"""
    files = {"file": ("test_questions.csv", io.BytesIO(csv_content.encode("utf-8")), "text/csv")}

    # 1. Preview
    res_prev = client.post(
        "/api/questions/import/preview",
        files=files,
        headers=examiner_auth_headers
    )
    assert res_prev.status_code == 200
    data = res_prev.json()
    assert data["total_detected"] == 4
    assert data["valid_count"] == 2
    assert data["invalid_count"] == 1
    assert data["duplicate_count"] == 1
    assert data["file_type"] == "CSV"

    # Verify row-level error on invalid item
    invalid_items = [q for q in data["questions"] if q["status"] == "INVALID"]
    assert len(invalid_items) == 1
    assert "Subject is required." in invalid_items[0]["errors"]

    # Verify duplicate warning on duplicate item
    dup_items = [q for q in data["questions"] if q["status"] == "DUPLICATE"]
    assert len(dup_items) == 1
    assert any("already exists" in w for w in dup_items[0]["warnings"])

    # 2. Confirm import (skip_invalid=True, allow_duplicates=False)
    res_conf = client.post(
        "/api/questions/import/confirm",
        json={
            "questions": data["questions"],
            "skip_invalid": True,
            "allow_duplicates": False
        },
        headers=examiner_auth_headers
    )
    assert res_conf.status_code == 200
    conf_data = res_conf.json()
    assert conf_data["imported_count"] == 2
    assert conf_data["skipped_invalid_count"] == 1
    assert conf_data["skipped_duplicate_count"] == 1

def test_excel_preview_and_confirm(client, examiner_auth_headers):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "subject", "topic", "question_text", "question_type", "difficulty", "marks", "negative_marks",
        "option_1", "option_1_correct", "option_2", "option_2_correct", "explanation"
    ])
    ws.append([
        "Biology", "Genetics", "What is the double-helix molecule carrying genetic instructions?", "MCQ", "EASY", 2.0, 0.5,
        "DNA", "TRUE", "RNA", "FALSE", "DNA encodes genetic information in living organisms."
    ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    files = {"file": ("test_genetics.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    res_prev = client.post(
        "/api/questions/import/preview",
        files=files,
        headers=examiner_auth_headers
    )
    assert res_prev.status_code == 200
    data = res_prev.json()
    assert data["file_type"] == "EXCEL"
    assert data["valid_count"] == 1
    assert data["questions"][0]["subject"] == "Biology"
    assert data["questions"][0]["question_text"] == "What is the double-helix molecule carrying genetic instructions?"

    # Confirm
    res_conf = client.post(
        "/api/questions/import/confirm",
        json={"questions": data["questions"], "skip_invalid": True, "allow_duplicates": False},
        headers=examiner_auth_headers
    )
    assert res_conf.status_code == 200
    assert res_conf.json()["imported_count"] == 1

def test_json_preview_and_confirm(client, examiner_auth_headers):
    json_data = [
        {
            "subject": "Astronomy",
            "topic": "Solar System",
            "question_text": "Which planet is known as the Red Planet?",
            "question_type": "MCQ",
            "difficulty": "EASY",
            "marks": 1.5,
            "negative_marks": 0.25,
            "options": [
                {"option_text": "Mars", "is_correct": True},
                {"option_text": "Venus", "is_correct": False},
                {"option_text": "Jupiter", "is_correct": False}
            ],
            "explanation": "Mars appears reddish due to iron oxide on its surface."
        }
    ]
    buf = io.BytesIO(json.dumps(json_data).encode("utf-8"))
    files = {"file": ("test_astronomy.json", buf, "application/json")}

    res_prev = client.post(
        "/api/questions/import/preview",
        files=files,
        headers=examiner_auth_headers
    )
    assert res_prev.status_code == 200
    data = res_prev.json()
    assert data["file_type"] == "JSON"
    assert data["valid_count"] == 1
    assert data["questions"][0]["subject"] == "Astronomy"

    res_conf = client.post(
        "/api/questions/import/confirm",
        json={"questions": data["questions"], "skip_invalid": True, "allow_duplicates": False},
        headers=examiner_auth_headers
    )
    assert res_conf.status_code == 200
    assert res_conf.json()["imported_count"] == 1

def test_docx_preview(client, examiner_auth_headers):
    doc = docx.Document()
    doc.add_paragraph("1. What is the powerhouse of the cell?")
    doc.add_paragraph("Subject: Biology")
    doc.add_paragraph("Type: MCQ")
    doc.add_paragraph("Difficulty: EASY")
    doc.add_paragraph("Marks: 2")
    doc.add_paragraph("A) Mitochondria [CORRECT]")
    doc.add_paragraph("B) Ribosome")
    doc.add_paragraph("C) Nucleus")
    doc.add_paragraph("Explanation: Mitochondria generates cellular energy via ATP.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    files = {"file": ("biology_cell.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    res = client.post(
        "/api/questions/import/preview",
        files=files,
        headers=examiner_auth_headers
    )
    assert res.status_code == 200
    data = res.json()
    assert data["file_type"] == "DOCX"
    assert data["total_detected"] >= 1
    q = data["questions"][0]
    assert q["subject"] == "Biology"
    assert "powerhouse of the cell" in q["question_text"]
    assert any(o["option_text"] == "Mitochondria" and o["is_correct"] for o in q["options"])

def test_pdf_preview(client, examiner_auth_headers):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("1. Which gas is most abundant in Earth's atmosphere?", styles["Normal"]),
        Paragraph("Subject: Earth Science", styles["Normal"]),
        Paragraph("Type: MCQ", styles["Normal"]),
        Paragraph("A) Nitrogen [CORRECT]", styles["Normal"]),
        Paragraph("B) Oxygen", styles["Normal"]),
        Paragraph("C) Argon", styles["Normal"]),
        Paragraph("Explanation: Nitrogen makes up approximately 78% of Earth's atmosphere.", styles["Normal"])
    ]
    doc.build(story)
    buf.seek(0)

    files = {"file": ("earth_science.pdf", buf, "application/pdf")}

    res = client.post(
        "/api/questions/import/preview",
        files=files,
        headers=examiner_auth_headers
    )
    assert res.status_code == 200
    data = res.json()
    assert data["file_type"] == "PDF"
    assert data["total_detected"] >= 1
    q = data["questions"][0]
    assert q["subject"] == "Earth Science"
    assert any("Nitrogen" in o["option_text"] for o in q["options"])

def test_export_all_formats(client, examiner_auth_headers):
    # Test Excel Export
    res_xl = client.get("/api/questions/export/excel", headers=examiner_auth_headers)
    assert res_xl.status_code == 200
    assert "spreadsheetml" in res_xl.headers["content-type"]
    assert len(res_xl.content) > 100

    # Test DOCX Export
    res_doc = client.get("/api/questions/export/docx", headers=examiner_auth_headers)
    assert res_doc.status_code == 200
    assert "wordprocessingml" in res_doc.headers["content-type"]
    assert len(res_doc.content) > 100

    # Test JSON Export
    res_json = client.get("/api/questions/export/json", headers=examiner_auth_headers)
    assert res_json.status_code == 200
    assert "application/json" in res_json.headers["content-type"]
    parsed = res_json.json()
    assert isinstance(parsed, list)

def test_templates_all_formats(client, examiner_auth_headers):
    res_xl = client.get("/api/questions/template/excel", headers=examiner_auth_headers)
    assert res_xl.status_code == 200
    assert "spreadsheetml" in res_xl.headers["content-type"]

    res_json = client.get("/api/questions/template/json", headers=examiner_auth_headers)
    assert res_json.status_code == 200
    assert "application/json" in res_json.headers["content-type"]

    res_docx = client.get("/api/questions/template/docx", headers=examiner_auth_headers)
    assert res_docx.status_code == 200
    assert "wordprocessingml" in res_docx.headers["content-type"]

def test_rbac_student_rejected(client, student_auth_headers):
    # Students cannot preview or confirm question import
    res = client.post(
        "/api/questions/import/preview",
        files={"file": ("test.csv", io.BytesIO(b"subject,question_text\nMath,Sample?"), "text/csv")},
        headers=student_auth_headers
    )
    assert res.status_code == 403

    res_conf = client.post(
        "/api/questions/import/confirm",
        json={"questions": [], "skip_invalid": True, "allow_duplicates": False},
        headers=student_auth_headers
    )
    assert res_conf.status_code == 403
